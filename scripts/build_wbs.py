"""Generate docs/03-wbs.xlsx (WBS + milestone + approval matrix + asumsi/risiko).

Jalankan: uv run python scripts/build_wbs.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "03-wbs.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
LEVEL_FILL = {
    1: PatternFill("solid", fgColor="DBEAFE"),
    2: PatternFill("solid", fgColor="EEF2FF"),
    3: PatternFill("solid", fgColor="F8FAFC"),
}
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# (kode, level, modul, fitur, sub-fitur/deliverable, prioritas, role, dependensi,
#  acceptance, estimasi, status)
WBS = [
    ("1", 1, "Foundation", "Setup & arsitektur", "", "Wajib", "All", "-", "", "", ""),
    (
        "1.1",
        2,
        "Foundation",
        "Project setup",
        "uv + Flask app factory + config + .env",
        "Wajib",
        "Dev",
        "-",
        "`uv sync` dan `flask run` berhasil",
        0.5,
        "Selesai",
    ),
    (
        "1.2",
        2,
        "Foundation",
        "Database",
        "SQLAlchemy models + init-db + seed deterministik",
        "Wajib",
        "Dev",
        "1.1",
        "`flask init-db && flask seed` menghasilkan data demo",
        1,
        "Selesai",
    ),
    (
        "1.3",
        2,
        "Foundation",
        "Autentikasi",
        "Login/logout session, halaman login SPA, 7 akun demo",
        "Wajib",
        "All",
        "1.1",
        "Setiap role bisa login dan melihat menu sesuai hak",
        0.5,
        "Selesai",
    ),
    (
        "1.4",
        2,
        "Foundation",
        "RBAC",
        "Decorator role + filter data per role/departemen",
        "Wajib",
        "All",
        "1.3",
        "Requester tidak bisa approve; officer tidak bisa verifikasi invoice",
        0.5,
        "Selesai",
    ),
    (
        "1.5",
        2,
        "Foundation",
        "Penomoran dokumen",
        "Sequence per tipe per bulan dalam transaksi DB",
        "Wajib",
        "System",
        "1.2",
        "Nomor unik berurutan, aman dari duplikasi",
        0.25,
        "Selesai",
    ),
    (
        "1.6",
        2,
        "Foundation",
        "Audit trail",
        "Activity log append-only + timeline per dokumen",
        "Wajib",
        "System",
        "1.2",
        "Semua transisi status tercatat dengan aktor & waktu",
        0.5,
        "Selesai",
    ),
    ("2", 1, "Master Data", "Referensi", "", "Wajib", "All", "", "", "", ""),
    (
        "2.1",
        2,
        "Master Data",
        "Departemen",
        "Data departemen + relasi user & budget",
        "Wajib",
        "Admin",
        "1.2",
        "Departemen tampil sebagai pilihan PR",
        0.25,
        "Selesai",
    ),
    (
        "2.2",
        2,
        "Master Data",
        "Budget",
        "Alokasi, committed, spent, available per tahun fiskal",
        "Wajib",
        "Finance",
        "2.1",
        "Available = allocated - committed - spent",
        0.5,
        "Selesai",
    ),
    (
        "2.3",
        2,
        "Master Data",
        "Vendor",
        "Data legal, kontak, kategori, status aktif/nonaktif, rating",
        "Wajib",
        "Procurement",
        "1.2",
        "Vendor blacklist tidak bisa diundang RFQ",
        0.5,
        "Selesai",
    ),
    (
        "2.4",
        2,
        "Master Data",
        "Item & UoM",
        "Katalog item, kategori, satuan, estimasi harga",
        "Wajib",
        "Procurement",
        "1.2",
        "Item bisa dipilih di form PR",
        0.5,
        "Selesai",
    ),
    ("3", 1, "Purchase Request", "Permintaan pengadaan", "", "Wajib", "Requester", "", "", "", ""),
    (
        "3.1",
        2,
        "Purchase Request",
        "Draft PR",
        "Form header + multi item + validasi budget departemen",
        "Wajib",
        "Requester",
        "2.2, 2.4",
        "PR tersimpan dengan total terhitung server",
        1,
        "Selesai",
    ),
    (
        "3.2",
        2,
        "Purchase Request",
        "Submit & revisi",
        "Submit, reject, revisi ke DRAFT, pembatalan",
        "Wajib",
        "Requester",
        "3.1",
        "Status mengikuti state machine; revisi menghapus approval lama",
        0.5,
        "Selesai",
    ),
    (
        "3.3",
        2,
        "Purchase Request",
        "Daftar & filter",
        "Tabel PR dengan filter status dan pencarian nomor/judul",
        "Wajib",
        "All",
        "3.1",
        "Role hanya melihat PR yang relevan",
        0.5,
        "Selesai",
    ),
    (
        "3.4",
        2,
        "Purchase Request",
        "Detail & timeline",
        "Item, budget, approval steps, activity log",
        "Wajib",
        "All",
        "1.6",
        "Riwayat lengkap terlihat di satu halaman",
        0.5,
        "Selesai",
    ),
    ("4", 1, "Approval", "Persetujuan berjenjang", "", "Wajib", "Approver", "", "", "", ""),
    (
        "4.1",
        2,
        "Approval",
        "Approval matrix",
        "Rule per doc type & nominal, seed 3 tier x 2 dokumen",
        "Wajib",
        "Admin",
        "1.2",
        "Tier benar terpilih berdasarkan nilai dokumen",
        0.5,
        "Selesai",
    ),
    (
        "4.2",
        2,
        "Approval",
        "Inbox approval",
        "Daftar tugas approval sesuai role, approve/reject + catatan",
        "Wajib",
        "Approver",
        "4.1",
        "Hanya step milik role yang bisa diproses",
        0.75,
        "Selesai",
    ),
    (
        "4.3",
        2,
        "Approval",
        "Commit budget",
        "Commit saat PR approved penuh; held oleh PR, bukan PO",
        "Wajib",
        "Finance",
        "4.1, 2.2",
        "Budget tidak bisa minus; commitment dirilis saat PR dibatalkan/ditutup",
        0.5,
        "Selesai",
    ),
    ("5", 1, "Sourcing", "RFQ & quotation", "", "Wajib", "Procurement", "", "", "", ""),
    (
        "5.1",
        2,
        "Sourcing",
        "RFQ",
        "Buat RFQ dari PR approved, pilih metode, undang vendor",
        "Wajib",
        "Officer",
        "3.2, 2.3",
        "RFQ minimal 1 vendor; metode sesuai nominal",
        0.75,
        "Selesai",
    ),
    (
        "5.2",
        2,
        "Sourcing",
        "Quotation",
        "Input quotation per vendor + item price, total server-side",
        "Wajib",
        "Officer",
        "5.1",
        "Total quotation konsisten dengan item",
        0.75,
        "Selesai",
    ),
    (
        "5.3",
        2,
        "Sourcing",
        "Comparison matrix",
        "Perbandingan harga, lead time, garansi, dan masa berlaku quotation",
        "Wajib",
        "Manager",
        "5.2",
        "Matrix menampilkan semua quotation & termurah ditandai",
        0.5,
        "Selesai",
    ),
    (
        "5.4",
        2,
        "Sourcing",
        "Award",
        "Tetapkan pemenang, minimal 2 quotation",
        "Wajib",
        "Manager",
        "5.2",
        "Hanya satu quotation SELECTED; RFQ menjadi AWARDED",
        0.5,
        "Selesai",
    ),
    ("6", 1, "Purchase Order", "PO", "", "Wajib", "Procurement", "", "", "", ""),
    (
        "6.1",
        2,
        "Purchase Order",
        "Generate PO",
        "Dari quotation pemenang, harga & vendor terkunci",
        "Wajib",
        "Officer",
        "5.4",
        "Satu quotation hanya satu PO",
        0.5,
        "Selesai",
    ),
    (
        "6.2",
        2,
        "Purchase Order",
        "Approval PO",
        "Matrix terpisah dari PR",
        "Wajib",
        "Manager/Finance/Director",
        "6.1, 4.1",
        "PO tidak bisa issued sebelum approved",
        0.5,
        "Selesai",
    ),
    (
        "6.3",
        2,
        "Purchase Order",
        "Issue & cetak",
        "Kirim ke vendor, halaman cetak PO",
        "Wajib",
        "Officer",
        "6.2",
        "Status ISSUED dan dokumen PO dapat dicetak",
        0.5,
        "Selesai",
    ),
    (
        "6.4",
        2,
        "Purchase Order",
        "Pembatalan",
        "Cancel dengan alasan; commitment tetap di PR sampai PR ditutup/dibatalkan",
        "Wajib",
        "Manager",
        "6.1",
        "PO dengan GRN tidak dapat dibatalkan",
        0.25,
        "Selesai",
    ),
    (
        "6.5",
        2,
        "Purchase Order",
        "Penutupan PO",
        "Tutup PO sisa, auto-close PR, rilis sisa commitment",
        "Wajib",
        "Manager",
        "6.4, 7.5",
        "PR otomatis CLOSED setelah semua PO selesai; sisa commitment dirilis",
        0.25,
        "Selesai",
    ),
    ("7", 1, "Receiving", "Penerimaan barang/jasa", "", "Wajib", "Warehouse", "", "", "", ""),
    (
        "7.1",
        2,
        "Receiving",
        "GRN",
        "Form per baris PO, qty diterima/ditolak, kondisi",
        "Wajib",
        "Warehouse",
        "6.3",
        "Qty kumulatif tidak melebihi PO",
        0.75,
        "Selesai",
    ),
    (
        "7.2",
        2,
        "Receiving",
        "Partial receipt",
        "Status PO partial/received otomatis",
        "Wajib",
        "Warehouse",
        "7.1",
        "Dua GRN parsial menghasilkan status RECEIVED",
        0.5,
        "Selesai",
    ),
    (
        "7.3",
        2,
        "Receiving",
        "Discrepancy",
        "Alasan rusak/salah spek/kurang, catatan retur",
        "Wajib",
        "Warehouse",
        "7.1",
        "Qty ditolak tercatat dan tidak menambah received",
        0.5,
        "Selesai",
    ),
    (
        "7.4",
        2,
        "Receiving",
        "BAST jasa",
        "Mode penerimaan jasa tanpa qty fisik",
        "Wajib",
        "Warehouse",
        "7.1",
        "Jasa dapat diselesaikan tanpa discrepancy",
        0.25,
        "Selesai",
    ),
    (
        "7.5",
        2,
        "Receiving",
        "Realisasi budget",
        "Pindahkan committed ke spent saat GRN posted",
        "Wajib",
        "Finance",
        "7.1, 4.3",
        "Spent bertambah sebesar nilai diterima",
        0.5,
        "Selesai",
    ),
    (
        "7.6",
        2,
        "Receiving",
        "Posting & pembatalan GRN",
        "GRN tersimpan sebagai DRAFT, diposting/dibatalkan lewat daftar GRN",
        "Wajib",
        "Warehouse",
        "7.1, 7.5",
        "Hanya GRN POSTED yang dihitung pada penerimaan dan 3-way match",
        0.25,
        "Selesai",
    ),
    ("8", 1, "Invoice", "Verifikasi invoice", "", "Wajib", "Finance", "", "", "", ""),
    (
        "8.1",
        2,
        "Invoice",
        "Input invoice",
        "Nomor, tanggal, nilai, vendor dari PO",
        "Wajib",
        "Officer",
        "7.1",
        "Invoice terkait PO dan tidak duplikat nomornya",
        0.5,
        "Selesai",
    ),
    (
        "8.2",
        2,
        "Invoice",
        "3-way match",
        "Cek qty diterima, nilai PO, toleransi 2%",
        "Wajib",
        "Finance",
        "8.1",
        "Match gagal menghasilkan DISPUTED dengan alasan",
        0.75,
        "Selesai",
    ),
    (
        "8.3",
        2,
        "Invoice",
        "Verifikasi & handoff",
        "VERIFIED lalu HANDED_OFF ke finance app",
        "Wajib",
        "Finance",
        "8.2",
        "Hanya invoice MATCHED yang bisa diverifikasi",
        0.5,
        "Selesai",
    ),
    (
        "8.4",
        2,
        "Invoice",
        "Cek ulang match",
        "Hitung ulang 3-way match setelah GRN baru diposting",
        "Wajib",
        "Officer",
        "8.2, 7.6",
        "Invoice DISPUTED dapat berubah menjadi MATCHED setelah penerimaan baru",
        0.25,
        "Selesai",
    ),
    ("9", 1, "Reporting", "Dashboard & laporan", "", "Wajib", "All", "", "", "", ""),
    (
        "9.1",
        2,
        "Reporting",
        "Dashboard",
        "Kartu status pipeline, dokumen terbaru, budget terpakai",
        "Wajib",
        "All",
        "3-8",
        "Angka konsisten dengan data seed",
        0.75,
        "Selesai",
    ),
    (
        "9.2",
        2,
        "Reporting",
        "Spend report",
        "Spend per departemen & kategori, cycle time PR→PO→GRN",
        "Wajib",
        "Management",
        "7.5",
        "Filter periode bekerja",
        0.75,
        "Selesai",
    ),
    (
        "9.3",
        2,
        "Reporting",
        "Export",
        "Export CSV daftar dokumen",
        "Opsional",
        "Management",
        "9.2",
        "CSV terbuka di Excel",
        0.25,
        "Backlog",
    ),
    ("10", 1, "Notification", "Notifikasi", "", "Wajib", "All", "", "", "", ""),
    (
        "10.1",
        2,
        "Notification",
        "In-app notification",
        "Bell + daftar notifikasi saat ada approval/status berubah",
        "Wajib",
        "All",
        "1.6",
        "Notifikasi muncul untuk user terkait",
        0.5,
        "Selesai",
    ),
    (
        "10.2",
        2,
        "Notification",
        "Realtime SSE",
        "Push notifikasi & chat internal (interpretasi fitur brief)",
        "Opsional",
        "All",
        "10.1",
        "Reconnect otomatis, pesan tersimpan",
        1,
        "Backlog",
    ),
    ("11", 1, "Delivery", "Dokumen & kualitas", "", "Wajib", "All", "", "", "", ""),
    (
        "11.1",
        2,
        "Delivery",
        "Riset proses bisnis",
        "docs/01-business-process.md",
        "Wajib",
        "All",
        "-",
        "Scope, role, aturan, edge case terdokumentasi",
        0.5,
        "Selesai",
    ),
    (
        "11.2",
        2,
        "Delivery",
        "BPMN",
        "Mermaid swimlane + flowchart + state + PDF",
        "Wajib",
        "All",
        "11.1",
        "PDF sesuai brief 1a",
        0.75,
        "Selesai",
    ),
    (
        "11.3",
        2,
        "Delivery",
        "WBS",
        "docs/03-wbs.xlsx (file ini)",
        "Wajib",
        "All",
        "11.1",
        "Sesuai brief 1b (Excel)",
        0.25,
        "Selesai",
    ),
    (
        "11.4",
        2,
        "Delivery",
        "ERD & arsitektur",
        "docs/04-er-model.md",
        "Wajib",
        "Dev",
        "1.2",
        "ERD sesuai implementasi",
        0.25,
        "Selesai",
    ),
    (
        "11.5",
        2,
        "Delivery",
        "Test otomatis",
        "pytest: matrix, state, budget, match, numbering",
        "Wajib",
        "Dev",
        "3-8",
        "`pytest` hijau",
        1,
        "Selesai",
    ),
    (
        "11.6",
        2,
        "Delivery",
        "README & demo",
        "Setup, akun demo, screenshot, known limitation",
        "Wajib",
        "All",
        "semua",
        "Reviewer bisa menjalankan dari clean clone",
        0.5,
        "Selesai",
    ),
    (
        "11.7",
        2,
        "Delivery",
        "AI/Computer Vision",
        "OCR surat jalan untuk prefill GRN",
        "Opsional",
        "Dev",
        "7.1",
        "Draft GRN terisi otomatis",
        2,
        "Backlog",
    ),
]

MILESTONES = [
    (
        "M0",
        "Analysis & foundation",
        "docs riset, BPMN + PDF, WBS XLSX, ERD, model DB, seed",
        "Semua deliverable soal 1 selesai",
        "Selesai",
    ),
    (
        "M1",
        "Core flow",
        "Auth/RBAC, master data, PR, approval engine, audit",
        "PR diajukan dan di-approve lintas role",
        "Selesai",
    ),
    (
        "M2",
        "Sourcing & PO",
        "RFQ, quotation, comparison, award, PO, cetak",
        "PR approved menghasilkan PO issued",
        "Selesai",
    ),
    (
        "M3",
        "Receiving & invoice",
        "GRN partial/discrepancy, BAST, 3-way match, handoff",
        "PO diterima dan invoice terverifikasi",
        "Selesai",
    ),
    (
        "M4",
        "Frontend & delivery",
        "14 layar SPA, build, README, screenshot, tests",
        "Demo end-to-end dari UI",
        "Selesai",
    ),
    (
        "M5",
        "Backlog",
        "Realtime chat/SSE, AI OCR GRN, portal vendor, ekspor CSV",
        "Opsional, tidak menghambat demo",
        "Backlog",
    ),
]

MATRIX = [
    ("Purchase Request", "0", "10000000", 1, "Dept Head"),
    ("Purchase Request", "0", "10000000", 2, "Finance"),
    ("Purchase Request", "10000001", "100000000", 1, "Dept Head"),
    ("Purchase Request", "10000001", "100000000", 2, "Finance"),
    ("Purchase Request", "10000001", "100000000", 3, "Procurement Manager"),
    ("Purchase Request", "100000001", None, 1, "Dept Head"),
    ("Purchase Request", "100000001", None, 2, "Finance"),
    ("Purchase Request", "100000001", None, 3, "Procurement Manager"),
    ("Purchase Request", "100000001", None, 4, "Director"),
    ("Purchase Order", "0", "10000000", 1, "Procurement Manager"),
    ("Purchase Order", "10000001", "100000000", 1, "Procurement Manager"),
    ("Purchase Order", "10000001", "100000000", 2, "Finance"),
    ("Purchase Order", "100000001", None, 1, "Procurement Manager"),
    ("Purchase Order", "100000001", None, 2, "Finance"),
    ("Purchase Order", "100000001", None, 3, "Director"),
]

RISKS = [
    (
        "Asumsi",
        "Fitur realtime (live chat/mobile) di brief Procurement adalah sisa template Helpdesk",
        "Tidak diimplementasikan; dicatat di README; notifikasi in-app tersedia",
        "Rendah",
    ),
    (
        "Asumsi",
        "Vendor tidak memiliki portal login",
        "Quotation diinput Procurement Officer",
        "Rendah",
    ),
    (
        "Asumsi",
        "Mata uang tunggal IDR, harga belum termasuk pajak",
        "Tidak ada field pajak/kurs",
        "Rendah",
    ),
    (
        "Risiko",
        "Scope P2P penuh terlalu besar untuk 1-2 hari",
        "Milestone M0-M4 dengan cut-line; fitur opsional dipindah ke backlog",
        "Sedang",
    ),
    (
        "Risiko",
        "Build SPA tidak di-commit",
        "README + `make build` + pesan error ramah bila dist belum ada",
        "Sedang",
    ),
    (
        "Risiko",
        "SQLite tidak mendukung row lock",
        "Transaksi tunggal + guard status; model kompatibel Postgres",
        "Rendah",
    ),
    (
        "Risiko",
        "Data demo tidak konsisten untuk laporan",
        "Seed deterministik dan idempoten (`flask seed-reset`)",
        "Rendah",
    ),
]


def style_header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def autosize(ws, widths):
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def build():
    wb = Workbook()

    ws = wb.active
    ws.title = "WBS"
    headers = [
        "Kode",
        "Level",
        "Modul",
        "Fitur",
        "Sub-fitur / Deliverable",
        "Prioritas",
        "Role",
        "Dependensi",
        "Kriteria Penerimaan",
        "Estimasi (hari-orang)",
        "Status",
    ]
    ws.append(headers)
    style_header(ws)

    for row in WBS:
        code, level, module, feature, sub, priority, role, deps, acceptance, estimate, status = row
        ws.append(
            [code, level, module, feature, sub, priority, role, deps, acceptance, estimate, status]
        )
        r = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=col)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=col >= 5)
        fill = LEVEL_FILL.get(level)
        if fill:
            for col in (1, 2, 3, 4):
                ws.cell(row=r, column=col).fill = fill
        if level == 1:
            for col in range(1, len(headers) + 1):
                ws.cell(row=r, column=col).font = Font(bold=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{ws.max_row}"
    autosize(ws, [10, 7, 20, 20, 52, 10, 18, 14, 52, 14, 11])

    total = sum(row[9] for row in WBS if isinstance(row[9], (int, float)))
    ws.append([])
    ws.append(["", "", "", "Total estimasi", "", "", "", "", "", round(total, 2), ""])

    ws2 = wb.create_sheet("Milestone")
    ws2.append(["Milestone", "Nama", "Lingkup", "Checkpoint", "Status"])
    style_header(ws2)
    for row in MILESTONES:
        ws2.append(list(row))
        for cell in ws2[ws2.max_row]:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws2.freeze_panes = "A2"
    autosize(ws2, [12, 24, 60, 45, 11])

    ws3 = wb.create_sheet("Approval Matrix")
    ws3.append(["Doc Type", "Min (IDR)", "Max (IDR)", "Level", "Role"])
    style_header(ws3)
    for row in MATRIX:
        ws3.append(list(row))
        for cell in ws3[ws3.max_row]:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top")
        ws3.cell(row=ws3.max_row, column=2).number_format = "#,##0"
        if row[2] is not None:
            ws3.cell(row=ws3.max_row, column=3).number_format = "#,##0"
    autosize(ws3, [22, 16, 16, 8, 24])

    ws4 = wb.create_sheet("Asumsi & Risiko")
    ws4.append(["Tipe", "Deskripsi", "Mitigasi / Keputusan", "Level"])
    style_header(ws4)
    for row in RISKS:
        ws4.append(list(row))
        for cell in ws4[ws4.max_row]:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    autosize(ws4, [10, 60, 65, 10])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT} ({len(WBS)} baris WBS)")


if __name__ == "__main__":
    build()
